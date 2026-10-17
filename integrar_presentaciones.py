"""Integra las preguntas en vivo a las presentaciones EyB 2027.
- Inserta una diapositiva con el QR antes de cada "Pregunta para la clase".
- Agrega el botón "Votar en vivo" (hipervínculo) en la diapositiva de la pregunta.
- Genera el banco de preguntas (CSV) para cargar en el panel.
Uso: python3 integrar.py carpeta_entrada carpeta_salida
"""
import sys, re, glob, os, io, csv, copy
import qrcode
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = "https://ppaccioretti.github.io/Preguntas-eyb/"
URL_ALUMNO = BASE
VERDE = RGBColor(0x1E, 0x4D, 0x3A)
MOSTAZA = RGBColor(0xE3, 0xA7, 0x2F)
CREMA = RGBColor(0xFB, 0xF0, 0xDC)
TEXTO = RGBColor(0x2E, 0x3A, 0x34)
GRIS = RGBColor(0x5F, 0x72, 0x68)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

# Ajuste de la presentación 01: la consigna original pide clasificar las 4 variables;
# para votar hace falta una sola respuesta correcta.
ESPECIAL_01 = {
    "texto": "¿Cuál de estas variables es cuantitativa continua?",
    "correcta": "C",
    "explicacion": "El contenido de proteína (%) se mide y puede tomar cualquier valor dentro de un rango. "
                   "A es cuantitativa discreta (se cuenta), B es cualitativa nominal y D es cualitativa ordinal.",
}
CORRECTA_MANUAL = {"04": "A"}
TEMAS = {"01": "Introducción y análisis exploratorio", "02": "Medidas resumen", "03": "Probabilidad",
         "04": "Distribución Normal y media muestral", "05": "Estimación", "06": "Contraste de hipótesis",
         "07": "ANAVA DCA", "08": "ANAVA DBCA", "09": "Experimentos factoriales", "10": "Regresión lineal",
         "11": "Correlación y Chi cuadrado"}


def qr_png(url):
    img = qrcode.make(url, border=1, box_size=12)
    b = io.BytesIO(); img.save(b, format="PNG"); b.seek(0)
    return b


def texto(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def caja(slide, x, y, w, h, txt, size, color, bold=False, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    lineas = txt if isinstance(txt, list) else [txt]
    for i, l in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i: p.space_before = Pt(10)
        r = p.add_run(); r.text = l
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = font
    return tb


def diapositiva_qr(pr, layout, n_clase):
    s = pr.slides.add_slide(layout)
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 100:
            ph.text_frame.text = "Respondé desde el celular"
    caja(s, 0.9, 1.75, 6.6, 3.2, [
        "1.  Escaneá el código con la cámara del celular.",
        "2.  Elegí tu comisión.",
        "3.  Dejá la página abierta: las preguntas aparecen solas.",
    ], 22, TEXTO)
    caja(s, 0.9, 4.75, 6.6, 1.0, ["No hace falta cuenta y las respuestas son anónimas.",
                                  URL_ALUMNO.replace("https://", "").rstrip("/")], 15, GRIS)
    marco = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(1.45), Inches(4.4), Inches(5.15))
    marco.adjustments[0] = 0.06
    marco.fill.solid(); marco.fill.fore_color.rgb = CREMA; marco.line.fill.background(); marco.shadow.inherit = False
    marco.name = "Marco QR"
    s.shapes.add_picture(qr_png(URL_ALUMNO), Inches(8.6), Inches(1.75), Inches(3.7), Inches(3.7)).name = "QR alumnos"
    caja(s, 8.25, 5.65, 4.4, 0.6, "Escaneá acá", 18, VERDE, bold=True, align=PP_ALIGN.CENTER, font="Cambria")
    s.notes_slide.notes_text_frame.text = (
        "Dar un minuto para que se conecten. En la vista de votación (botón Votar en vivo de la diapositiva siguiente) "
        "elegí la comisión arriba a la derecha antes de lanzar. Si alguien cierra la página, en la vista de votación "
        "aparece un QR chico mientras la votación está abierta.")
    return s


def boton_votar(slide, url):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.6), Inches(0.5), Inches(2.13), Inches(0.58))
    b.adjustments[0] = 0.25
    b.fill.solid(); b.fill.fore_color.rgb = VERDE; b.line.fill.background(); b.shadow.inherit = False
    b.name = "Boton Votar en vivo"
    tf = b.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "▶  Votar en vivo"
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BLANCO; r.font.name = "Calibri"
    b.click_action.hyperlink.address = url
    return b


def procesar(ruta_in, ruta_out):
    num = os.path.basename(ruta_in)[:2]
    tema = re.sub(r"^\d+_", "", os.path.splitext(os.path.basename(ruta_in))[0]).replace("_", " ")
    pr = Presentation(ruta_in)
    slides = list(pr.slides)
    idx = next(i for i, s in enumerate(slides) if any(texto(sh) == "Pregunta para la clase" for sh in s.shapes))
    q = slides[idx]
    cajas = [sh for sh in q.shapes if sh.has_text_frame and texto(sh) and not sh.is_placeholder]
    opciones = sorted([sh for sh in cajas if 2.2 < sh.left / 914400 < 2.45], key=lambda s: s.top)
    y_op = min(o.top for o in opciones)
    enunciado = sorted([sh for sh in cajas if 1.5 < sh.left / 914400 < 1.7 and sh.top < y_op], key=lambda s: s.top)
    notas = q.notes_slide.notes_text_frame.text if q.has_notes_slide else ""

    qid = f"C{num}"
    url = f"{BASE}docente.html?q={qid}"
    if num == "01":
        tf = enunciado[0].text_frame
        runs = tf.paragraphs[0].runs
        runs[0].text = ESPECIAL_01["texto"]
        for r in runs[1:]: r.text = ""
        txt_preg, correcta, expl = ESPECIAL_01["texto"], ESPECIAL_01["correcta"], ESPECIAL_01["explicacion"]
        q.notes_slide.notes_text_frame.text = (f"Respuesta C. {expl}\nPara ampliar en voz alta: "
                                               f"clasificar las cuatro. {notas}")
    else:
        txt_preg = "\n".join(texto(e) for e in enunciado)
        m = re.match(r"\s*Respuesta\s+([A-E])\s*[\.:]?\s*", notas)
        correcta = m.group(1) if m else CORRECTA_MANUAL[num]
        expl = notas[m.end():].strip() if m else notas.strip()

    boton_votar(q, url)
    nf = q.notes_slide.notes_text_frame
    nf.text = nf.text + f"\n\nVotar en vivo: pregunta {qid} del banco."

    # diapositiva del QR antes de la pregunta
    nueva = diapositiva_qr(pr, q.slide_layout, num)
    lista = pr.slides._sldIdLst
    el = lista[-1]; lista.remove(el); lista.insert(idx, el)
    pr.save(ruta_out)
    return {"id": qid, "tema": f"Clase {num}: {TEMAS.get(num, tema)}", "pregunta": txt_preg,
            **{L: (texto(o) if i < len(opciones) else "") for i, (L, o) in enumerate(zip("ABCDE", opciones + [None] * 5))},
            "correcta": correcta, "explicacion": expl, "revelar_celular": "si", "imagen": ""}


if __name__ == "__main__":
    ent, sal = sys.argv[1], sys.argv[2]
    os.makedirs(sal, exist_ok=True)
    filas = []
    for f in sorted(glob.glob(os.path.join(ent, "*.pptx"))):
        filas.append(procesar(f, os.path.join(sal, os.path.basename(f))))
        print("ok", os.path.basename(f), filas[-1]["id"], filas[-1]["correcta"])
    with open(os.path.join(sal, "banco_presentaciones_2027.csv"), "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["id", "tema", "pregunta", "A", "B", "C", "D", "E", "correcta", "explicacion", "revelar_celular", "imagen"], delimiter=";")
        w.writeheader(); w.writerows(filas)
